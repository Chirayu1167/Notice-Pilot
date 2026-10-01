"""Generate SYNTHETIC DEMO return fixtures (NOT real GST data). Run from backend/."""
import csv, os

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures", "returns")
os.makedirs(BASE, exist_ok=True)
TAG = "SYNTHETIC DEMO DATA \u2014 NOT REAL GST DATA"

def write_csv(name, headers, rows):
    path = os.path.join(BASE, name)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(headers)
        w.writerows(rows)
    print("wrote", path)

def write_xlsx(name, sheets):
    import openpyxl
    path = os.path.join(BASE, name)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.create_sheet("README")
    wb["README"]["A1"] = TAG + " \u2014 every row in this workbook is fictitious."
    for sheet_name, (headers, rows) in sheets.items():
        ws = wb.create_sheet(sheet_name)
        ws.append(headers)
        for r in rows:
            ws.append(r)
    wb.save(path)
    print("wrote", path)

G1H = ["GSTIN/UIN of Recipient", "Invoice Number", "Invoice date", "Invoice Value",
       "Place Of Supply", "Rate", "Taxable Value", "IGST", "CGST", "SGST", "Cess", "Remarks"]
G3H = ["Nature of supplies", "Total Taxable Value", "IGST", "CGST", "SGST", "Cess", "Remarks"]
G2H = ["Supplier GSTIN", "Invoice Number", "Invoice Date", "Taxable Value",
       "IGST", "CGST", "SGST", "Cess", "ITC Available", "Remarks"]

# --- DRC-01B matched: GSTR-1 total 54000, GSTR-3B total 45000, diff 9000 ---
g1_matched = [
    ["27SYNTH0001A1Z5", "INV-001", "05/04/2024", "118000", "27-Maharashtra", "18", "100000", "18000", "0", "0", "0", TAG],
    ["27SYNTH0001A1Z5", "INV-002", "12/04/2024", "118000", "27-Maharashtra", "18", "100000", "18000", "0", "0", "0", TAG],
    ["27SYNTH0001A1Z5", "INV-003", "20/04/2024", "118000", "27-Maharashtra", "18", "100000", "0", "9000", "9000", "0", TAG],
]
g3b_matched = [
    ["Outward taxable supplies (other than zero rated, nil rated and exempted)", "250000", "45000", "0", "0", "0", TAG],
]
write_csv("gstr1_matched_2024-04.csv", G1H, g1_matched)
write_csv("gstr3b_matched_2024-04.csv", G3H, g3b_matched)
write_xlsx("gstr1_matched_2024-04.xlsx", {"B2B Invoices": (G1H, g1_matched)})
write_xlsx("gstr3b_matched_2024-04.xlsx", {"Table 3.1": (G3H, g3b_matched)})

# --- DRC-01C matched: GSTR-2B ITC 30000, GSTR-3B ITC 35000, diff 5000 ---
g2_matched = [
    ["29SUPPL0001B1Z2", "SUP-101", "03/04/2024", "100000", "10000", "0", "0", "0", "10000", TAG],
    ["29SUPPL0002C1Z3", "SUP-102", "11/04/2024", "100000", "0", "5000", "5000", "0", "10000", TAG],
    ["29SUPPL0003D1Z4", "SUP-103", "19/04/2024", "100000", "10000", "0", "0", "0", "10000", TAG],
]
g3b_itc = [
    ["Table 4 - Eligible ITC claimed", "200000", "20000", "7500", "7500", "0", TAG],
]
write_csv("gstr2b_matched_2024-04.csv", G2H, g2_matched)
write_csv("gstr3b_itc_2024-04.csv", G3H, g3b_itc)
write_xlsx("gstr2b_matched_2024-04.xlsx", {"GSTR-2B ITC": (G2H, g2_matched)})

# --- Period mismatch: same GSTR-1 shape but May 2024 dates ---
g1_may = [r[:] for r in g1_matched]
for r in g1_may:
    r[2] = r[2].replace("/04/2024", "/05/2024")
write_csv("gstr1_may_2024-05.csv", G1H, g1_may)

# --- Duplicate invoice: INV-001 twice ---
g1_dup = g1_matched + [[
    "27SYNTH0001A1Z5", "INV-001", "05/04/2024", "118000", "27-Maharashtra", "18", "100000", "18000", "0", "0", "0", TAG,
]]
write_csv("gstr1_duplicate_2024-04.csv", G1H, g1_dup)

# --- Unsupported / arbitrary CSV ---
write_csv("arbitrary_expenses.csv", ["Date", "Category", "Spent", "Notes"],
          [["2024-04-01", "Office snacks", "500", TAG]])

# --- Empty CSV (headers only) ---
write_csv("empty_returns.csv", G1H, [])

# --- Numeric-format torture: commas, rupee symbols, blanks ---
write_csv("gstr1_formats_2024-04.csv", G1H, [
    ["27SYNTH0001A1Z5", "INV-101", "2024-04-05", "118000", "27-Maharashtra", "18", "\u20b91,00,000.00", "\u20b918,000", "", "", "", TAG],
    ["27SYNTH0001A1Z5", "inv-102 ", "05-04-2024", "59000", "27-Maharashtra", "18", "50000", "0", "4500", "4500", "", TAG],
])

with open(os.path.join(BASE, "README.md"), "w", encoding="utf-8") as f:
    f.write("# SYNTHETIC DEMO DATA \u2014 NOT REAL GST DATA\n\n"
            "Every file in this directory is fictitious test data for the Phase-2\n"
            "reconciliation engine. GSTINs, invoice numbers and amounts are invented\n"
            "and internally consistent. No real taxpayer data.\n")
print("done")
